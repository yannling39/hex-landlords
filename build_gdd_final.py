from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'GDD v1.2 开工最终版.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def borders(table, color='D9D9D9'):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders')
        tblPr.append(b)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn('w:' + edge)
        el = b.find(tag)
        if el is None:
            el = OxmlElement('w:' + edge)
            b.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), color)

def set_cell(cell, text, bold=False, white=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(9)
    if white:
        r.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, True, True)
        shade(table.rows[0].cells[i], '44546A')
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if row_index % 2 == 1:
                shade(cells[i], 'F2F5F8')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
for style_name, size in [('Normal', 10), ('Title', 20), ('Heading 1', 15), ('Heading 2', 12)]:
    style = doc.styles[style_name]
    style.font.name = 'Microsoft YaHei'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    style.font.size = Pt(size)
    if style_name != 'Normal':
        style.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('三人斗地主 roguelike 游戏 GDD v1.2 开工最终版')
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('标准斗地主底层规则 + 海克斯扩展接口 | 2026-09-25')
run.italic = True

doc.add_heading('1. 开工决策', level=1)
doc.add_paragraph('本版本是明日开始搭建底层逻辑的唯一实现依据。底层采用标准三人斗地主的牌组、发牌、叫地主、牌型和结算规则；海克斯保留为局间扩展接口，具体强化效果暂不冻结，也不阻塞斗地主核心开发。')
doc.add_paragraph('原 GDD v1.0 与 v1.1 中的“每人 18 张、Trick 最多 12 轮、三人轮流出牌”不再作为斗地主底层规则使用。它们可以作为后续实验玩法的参考，但不能与本版本的牌局状态混用。')

doc.add_heading('2. MVP 范围', level=1)
add_table(doc, ['范围项', 'MVP 决定', '暂不实现'], [
    ['平台', '桌面和移动浏览器的单机 Web', '账号、联网房间、云存档'],
    ['对局', '1 名人类玩家 + 2 个规则型 AI', '在线 PvP、三人同设备轮流操作'],
    ['牌局结构', '一个 Run 包含 3 手斗地主', '无限局、赛季和排行榜'],
    ['核心玩法', '发牌、叫地主、出牌、牌型比较、结算', '复杂 AI、春天/反春天'],
    ['海克斯', '保留数据结构和局间触发点，默认关闭具体效果', '具体强化池、数值和构筑规则'],
])

doc.add_heading('3. 牌组与发牌', level=1)
bullet(doc, '使用标准 54 张牌：四种花色的 3 到 A、2，共 52 张；小王和大王各 1 张。')
bullet(doc, '牌力顺序为 3、4、5、6、7、8、9、10、J、Q、K、A、2、小王、大王；花色不参与牌力比较。')
bullet(doc, '每手牌洗牌后发给三名玩家各 17 张，剩余 3 张作为底牌，底牌在叫地主完成后归地主。')
bullet(doc, '发牌顺序从本手的首位玩家开始，按 A、B、C 循环；首位玩家在三手 Run 中依次轮换，便于测试。')

doc.add_heading('4. 叫地主', level=1)
bullet(doc, '首位玩家先叫，随后按固定顺序轮流叫地主；每次可选择 1 分、2 分、3 分或不叫。')
bullet(doc, '叫分必须严格高于当前最高分；叫到 3 分立即确定该玩家为地主。')
bullet(doc, '三名玩家都不叫时，按同一套牌重新发牌，最多重发 2 次；第三次仍全不叫时，由首位玩家自动成为地主，底分为 1。')
bullet(doc, '地主获得 3 张底牌并明示给所有玩家；地主先出牌。两名未成为地主的玩家组成农民方。')
bullet(doc, 'MVP AI 叫分使用固定策略：有高价值牌型或大牌时倾向叫 2/3 分，否则不叫；策略可配置但不影响规则层。')

doc.add_heading('5. MVP 牌型与比较', level=1)
doc.add_paragraph('底层牌型识别必须输出牌型、主牌值、长度和附带牌数量，禁止由 UI 或 AI 重复实现比较逻辑。')
add_table(doc, ['牌型', '牌数要求', '比较方式'], [
    ['单张', '1', '比较单牌牌力'],
    ['对子', '2 张同点数', '比较对子点数'],
    ['三张', '3 张同点数', '比较三张点数'],
    ['三带一', '3 张同点数 + 1 单牌', '比较三张点数'],
    ['三带二', '3 张同点数 + 1 对子', '比较三张点数'],
    ['顺子', '至少 5 张连续单牌', '同长度比较最高牌；不得含 2 和王'],
    ['连对', '至少 3 个连续对子', '同长度比较最高对子；不得含 2 和王'],
    ['飞机带单', '至少 2 个连续三张 + 等量单牌', '同飞机长度比较最高三张；不得含 2 和王'],
    ['四带二', '4 张同点数 + 2 张单牌', '比较四张点数；两张带牌不参与比较'],
    ['炸弹', '4 张同点数', '可压过所有非炸弹；炸弹之间比较点数'],
    ['王炸', '小王 + 大王', '可压过所有牌型'],
])
bullet(doc, '顺子、连对和飞机的连续序列只允许 3 到 A，不允许出现 2、王或跨越 2 的序列。')
bullet(doc, '普通牌型只能压过同牌型、同长度且主牌值更高的牌；不同普通牌型不能互相比较。')
bullet(doc, '炸弹可压过所有普通牌型；更大的炸弹可压过更小的炸弹；王炸最大。')
bullet(doc, '四带二在 MVP 中只支持“带两张单牌”，不支持“四带两对”；飞机在 MVP 中只支持“带等量单牌”，不支持“带对子”。')
bullet(doc, '首轮出牌不需要压过上一手；非首轮必须出牌或 Pass。连续 Pass 后，由最后一次有效出牌者重新领出。')

doc.add_heading('6. 牌局状态机', level=1)
add_table(doc, ['状态', '进入条件', '系统/玩家动作', '离开条件'], [
    ['RUN_START', '点击开始 Run', '清空总分，设置手数和首位玩家', '进入 DEAL'],
    ['DEAL', '开始一手', '洗牌、发 17 张手牌、保留 3 张底牌', '进入 BID'],
    ['BID', '发牌完成', '玩家按顺序叫 1/2/3 分或不叫', '确定地主后进入 LANDLORD_REVEAL'],
    ['LANDLORD_REVEAL', '地主确定', '把底牌加入地主手牌并展示', '地主确认后进入 PLAY'],
    ['PLAY', '牌局未结束', '当前玩家提交合法牌型或 Pass', '有人手牌为空后进入 HAND_END'],
    ['HAND_END', '一方出完牌', '判断地主方胜负、炸弹倍数和本手得分', '海克斯开启则进入 HEX_DRAFT，否则进入 NEXT_HAND'],
    ['HEX_DRAFT', '局间扩展开启', '生成候选强化并选择；MVP 默认跳过', '选择完成后进入 NEXT_HAND'],
    ['NEXT_HAND', '本手结算完成', '轮换首位玩家和手数', '未完成 3 手则 DEAL，否则 RUN_END'],
    ['RUN_END', '完成 3 手', '累计分数并显示 Run 结果', '等待重新开始'],
])

doc.add_heading('7. 结算与倍数', level=1)
bullet(doc, '本手底分等于地主最终叫分，取值 1、2 或 3。')
bullet(doc, '每出现一次炸弹或王炸，倍数乘 2；MVP 不设置倍数上限，但必须使用整数状态并在 UI 显示。')
bullet(doc, '地主胜利：地主获得 2 × 底分 × 倍数；每名农民扣除底分 × 倍数。')
bullet(doc, '地主失败：地主扣除 2 × 底分 × 倍数；每名农民获得底分 × 倍数。')
bullet(doc, 'MVP 暂不实现春天、反春天、加倍、明牌、超级加倍和其他平台差异规则。')
bullet(doc, 'Run 结束后按三手累计分数最高者显示获胜方；同分显示平局。')

doc.add_heading('8. 海克斯扩展接口', level=1)
doc.add_paragraph('海克斯不再使用当前版本中未经验证的具体强化效果。底层只实现可扩展的数据接口和局间状态，不把未确定的数值或牌局改写规则写入斗地主核心。')
bullet(doc, '触发点：每手结算完成后、下一手发牌前；不在出牌过程中打断牌型比较。')
bullet(doc, '接口状态建议包含：hexEnabled、候选列表、当前选择玩家、已选强化 ID、剩余持续手数。')
bullet(doc, 'MVP 默认 hexEnabled=false；即使开关为 false，HAND_END 到 NEXT_HAND 的状态转移仍必须正常工作。')
bullet(doc, '后续设计需要单独确定强化的目标：影响叫分、牌型、倍数、AI 决策或资源成长；确定前不得直接修改牌型比较器。')
bullet(doc, '海克斯具体内容、候选池、刷新、升级、跨手保留和 UI 表现列入后续设计，不影响明日底层开发。')

doc.add_heading('9. AI 最小规格', level=1)
bullet(doc, 'AI 必须调用与玩家相同的合法出牌和牌型比较接口，不得自行判断牌型。')
bullet(doc, '叫地主采用可配置的简单评分策略；出牌优先使用能压过当前牌型的最小合法牌，没有合法牌时 Pass。')
bullet(doc, 'AI 在自己领出时优先出最小单张或最小对子；炸弹和王炸只在无法合理跟牌或牌局结束有利时使用。')
bullet(doc, 'AI 行动延迟 300 到 800 毫秒；超时自动执行默认动作，不能阻塞状态机。')

doc.add_heading('10. 底层模块边界', level=1)
add_table(doc, ['模块', '职责', '禁止承担的职责'], [
    ['Card / Deck', '牌面、牌组、洗牌、发牌', 'UI、AI 策略'],
    ['Pattern', '牌型识别、合法性、比较', '修改玩家分数'],
    ['Bid', '叫地主顺序、叫分和地主确定', '决定牌型比较'],
    ['Play', '出牌、Pass、牌权和回合', '渲染界面'],
    ['Scoring', '底分、倍数、地主/农民结算', '生成随机牌'],
    ['Game State', '状态机、事件和重开', '包含具体 UI 组件'],
    ['AI', '基于公开状态选择叫分和出牌', '绕过规则接口'],
    ['Hex Adapter', '局间候选和选择接口，默认关闭', '直接改写核心牌型规则'],
])

doc.add_heading('11. 7 天开工计划', level=1)
add_table(doc, ['日期', '任务', '验收产出'], [
    ['D1', '建立项目骨架；完成 Card、Deck、发牌和确定性随机种子', '54 张牌唯一性、17/17/17+3 发牌通过'],
    ['D2', '完成 Pattern 识别与比较：单张、对子、三张、三带、顺子、连对、炸弹、王炸', '牌型单元案例通过，非法牌型被拒绝'],
    ['D3', '完成 BID、地主确定、底牌归属和 PLAY 状态机', '可从发牌跑到一方出完牌'],
    ['D4', '完成 Scoring、倍数、手牌结束和 3 手 Run 流程', '可连续完成 3 手并得到稳定分数'],
    ['D5', '接入两个 AI 和超时处理；完成重开、异常状态恢复', '人类玩家可完成完整 Run'],
    ['D6', '接入基础 UI、牌型提示、牌权提示、底牌展示和结果页', '可演示版本，无卡死操作'],
    ['D7', '回归测试、修复 P0/P1、冻结接口、打包部署', 'MVP 构建包、测试记录和已知限制'],
])

doc.add_heading('12. 开工前验收标准', level=1)
for item in [
    '每手牌都能正确生成 54 张牌，三名玩家各 17 张，底牌 3 张且无重复。',
    '叫地主严格遵守顺序、叫分递增、全不叫重发和自动地主规则。',
    '地主拿到底牌并先出牌，农民双方身份在整手牌中保持一致。',
    '牌型识别、合法性和比较结果不依赖 UI 或 AI，且覆盖 MVP 牌型表。',
    '普通牌型、炸弹和王炸的压制关系符合第 5 节。',
    '有人手牌为空后立即结束本手，不再产生额外出牌操作。',
    '炸弹/王炸倍数和地主/农民得分符合第 7 节。',
    'hexEnabled=false 时不出现海克斯选择，但状态机仍能完成完整 Run。',
    '连续完成 3 个 Run 不出现牌数错误、非法出牌、分数回退或状态卡死。',
]: bullet(doc, item)

doc.add_heading('13. 后续讨论项', level=1)
doc.add_paragraph('海克斯具体强化、春天与反春天、完整斗地主牌型扩展、联网 PvP、复杂 AI、成长系统、商店、排行榜、音效和定制美术均在 MVP 核心稳定后单独立项。任何新规则必须同时更新 Pattern、Game State、Scoring、测试案例和验收标准。')

doc.add_heading('14. 版本决策记录', level=1)
doc.add_paragraph('v1.2 的核心判断是：明日先建立可验证的标准斗地主底层，再通过 Hex Adapter 接入局间随机强化。这样海克斯可以继续讨论和调整，而不会反复改动发牌、牌型、牌权和结算的基础模型。')

doc.save(OUT)
print(OUT)
